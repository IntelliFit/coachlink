from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
INPUT_MD = ROOT / "docs" / "CoachLink_商业计划书_v2_1.md"
OUTPUT_DOCX = ROOT / "output" / "docx" / "CoachLink_商业计划书_v2_1.docx"

FONT = "等线"
FONT_EN = "Aptos"
ACCENT = "C65F46"
ACCENT_DARK = "8F3F2E"
TEXT = "2B3340"
MUTED = "64707D"
SOFT_BG = "F6F2EE"
SOFT_BLUE = "EEF4F8"
GRID = "D9DEE3"


def set_paragraph_spacing(
    paragraph,
    *,
    before: float = 0,
    after: float = 0,
    line: float | None = None,
    left: float | None = None,
    right: float | None = None,
) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if line is not None:
        fmt.line_spacing = line
    if left is not None:
        fmt.left_indent = Cm(left)
    if right is not None:
        fmt.right_indent = Cm(right)


def normalize_text(text: str) -> str:
    replacements = {
        "\u2014": "-",
        "\u2013": "-",
        "\u2212": "-",
        "\u00a0": " ",
        "\ufe0f": "",
        "\u200b": "",
    }
    for src, dst in replacements.items():
        text = text.replace(src, dst)
    text = re.sub(r"[\U0001F300-\U0001FAFF]", "", text)
    text = text.replace("**", "").replace("__", "")
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = FONT_EN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(paragraph, size: float = 10.5, bold: bool = False, color: str = TEXT) -> None:
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRID) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def style_table(table, header: bool = True) -> None:
    table.autofit = True
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            if header and row_idx == 0:
                shade_cell(cell, ACCENT)
            else:
                shade_cell(cell, "FFFFFF" if row_idx % 2 else "FBFBFA")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=9,
                        bold=(header and row_idx == 0),
                        color=("FFFFFF" if header and row_idx == 0 else TEXT),
                    )


def set_doc_defaults(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.3)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_EN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.18

    for style_name, size, color in [
        ("Heading 1", 18, TEXT),
        ("Heading 2", 14, ACCENT_DARK),
        ("Heading 3", 12, TEXT),
    ]:
        style = styles[style_name]
        style.font.name = FONT_EN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        if style_name == "Heading 1":
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.line_spacing = 1.05
        elif style_name == "Heading 2":
            style.paragraph_format.space_before = Pt(8)
            style.paragraph_format.space_after = Pt(4)
            style.paragraph_format.line_spacing = 1.05
        else:
            style.paragraph_format.space_before = Pt(6)
            style.paragraph_format.space_after = Pt(3)
            style.paragraph_format.line_spacing = 1.05


def add_footer(section) -> None:
    section.different_first_page_header_footer = True
    first_footer = section.first_page_footer.paragraphs[0]
    first_footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(first_footer, before=0, after=0, line=1.0)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(footer, before=0, after=0, line=1.0)
    run = footer.add_run("CoachLink 商业计划书 v2.1")
    set_run_font(run, size=8, color=MUTED)


def add_cover(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(title, after=4, line=1.0)
    run = title.add_run("CoachLink 商业计划书 v2.1")
    set_run_font(run, size=24, bold=True, color=TEXT)

    subtitle = doc.add_paragraph()
    set_paragraph_spacing(subtitle, after=5, line=1.0)
    run = subtitle.add_run("线上自由教练的商业化操作系统")
    set_run_font(run, size=14, bold=True, color=ACCENT_DARK)

    body = doc.add_paragraph()
    set_paragraph_spacing(body, after=7, line=1.12)
    run = body.add_run(
        "从“力量训练工具”升级为“自由教练商业化基础设施”：不替教练抢流量，"
        "而是帮助教练把已有内容影响力沉淀为可订阅、可交付、可续费的线上训练服务。"
    )
    set_run_font(run, size=11, color=TEXT)

    metrics_title = doc.add_paragraph()
    set_paragraph_spacing(metrics_title, before=1, after=2, line=1.0)
    run = metrics_title.add_run("当前关键数据")
    set_run_font(run, size=10.5, bold=True, color=ACCENT_DARK)

    metrics = [
        ("注册用户", "1,296", "已批准教练", "15"),
        ("活跃教练", "3", "付费订阅用户", "14"),
        ("教练订阅关系", "126", "完成训练记录", "794"),
        ("完成整套计划", "5", "上传训练视频用户", "16"),
    ]

    for left_label, left_value, right_label, right_value in metrics:
        row = doc.add_paragraph()
        set_paragraph_spacing(row, after=1.2, line=1.0)
        r = row.add_run(left_label)
        set_run_font(r, size=10.3, bold=True, color=ACCENT_DARK)
        r = row.add_run(f"  {left_value}    |    ")
        set_run_font(r, size=10.3, color=TEXT)
        r = row.add_run(right_label)
        set_run_font(r, size=10.3, bold=True, color=ACCENT_DARK)
        r = row.add_run(f"  {right_value}")
        set_run_font(r, size=10.3, color=TEXT)

    quote = doc.add_paragraph()
    set_paragraph_spacing(quote, before=6, after=4, line=1.1, left=0.12, right=0.08)
    run = quote.add_run(
        "当前阶段判断：产品闭环已初步跑通，但增长瓶颈不在功能数量，而在教练从入驻到收款的首个成功路径。"
        "下一阶段应优先打造标杆教练、简化发布服务流程，并把交付数据转化为续费理由。"
    )
    set_run_font(run, size=10.5, color=ACCENT_DARK)

    meta = doc.add_paragraph()
    set_paragraph_spacing(meta, before=3, after=0, line=1.0)
    run = meta.add_run("版本：v2.1  |  日期：2026年4月  |  输出：可编辑 Word 版本")
    set_run_font(run, size=9, color=MUTED)


def strip_intro(lines: list[str]) -> list[str]:
    i = 0
    if lines and lines[0].startswith("# "):
        i = 1
    while i < len(lines) and (not lines[i].strip() or lines[i].startswith(">")):
        i += 1
    if i < len(lines) and lines[i].strip() == "---":
        i += 1
    while i < len(lines) and not lines[i].strip():
        i += 1
    return lines[i:]


def add_markdown_table(doc: Document, table_lines: list[str]) -> None:
    rows: list[list[str]] = []
    for line in table_lines:
        parts = [normalize_text(part) for part in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", part) for part in parts):
            continue
        rows.append(parts)
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c in range(col_count):
            cell = table.cell(r, c)
            cell.text = row[c] if c < len(row) else ""
    style_table(table)


def flush_paragraph(doc: Document, paragraph_lines: list[str]) -> None:
    if not paragraph_lines:
        return
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, after=3, line=1.16)
    run = paragraph.add_run(normalize_text(" ".join(paragraph_lines)))
    set_run_font(run, size=10.5, color=TEXT)
    paragraph_lines.clear()


def add_markdown_body(doc: Document, markdown: str) -> None:
    lines = strip_intro(markdown.splitlines())
    paragraph_lines: list[str] = []
    i = 0

    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            flush_paragraph(doc, paragraph_lines)
            i += 1
            continue

        if stripped == "---":
            flush_paragraph(doc, paragraph_lines)
            i += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph(doc, paragraph_lines)
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_markdown_table(doc, table_lines)
            continue

        if stripped.startswith("```"):
            flush_paragraph(doc, paragraph_lines)
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            paragraph = doc.add_paragraph()
            set_paragraph_spacing(paragraph, after=4, line=1.05)
            run = paragraph.add_run("\n".join(code_lines))
            set_run_font(run, size=9, color=TEXT)
            continue

        if stripped.startswith("## "):
            flush_paragraph(doc, paragraph_lines)
            paragraph = doc.add_paragraph(style="Heading 1")
            paragraph.add_run(normalize_text(stripped[3:]))
            set_paragraph_font(paragraph, size=18, bold=True, color=TEXT)
            i += 1
            continue

        if stripped.startswith("### "):
            flush_paragraph(doc, paragraph_lines)
            paragraph = doc.add_paragraph(style="Heading 2")
            paragraph.add_run(normalize_text(stripped[4:]))
            set_paragraph_font(paragraph, size=14, bold=True, color=ACCENT_DARK)
            i += 1
            continue

        if stripped.startswith("#### "):
            flush_paragraph(doc, paragraph_lines)
            paragraph = doc.add_paragraph(style="Heading 3")
            paragraph.add_run(normalize_text(stripped[5:]))
            set_paragraph_font(paragraph, size=12, bold=True, color=TEXT)
            i += 1
            continue

        if stripped.startswith(">"):
            flush_paragraph(doc, paragraph_lines)
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            paragraph = doc.add_paragraph()
            set_paragraph_spacing(paragraph, before=2, after=4, line=1.12, left=0.24)
            run = paragraph.add_run(normalize_text(" ".join(quote_lines)))
            set_run_font(run, size=10.5, color=ACCENT_DARK)
            continue

        bullet = re.match(r"^-\s+(.+)", stripped)
        if bullet:
            flush_paragraph(doc, paragraph_lines)
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.add_run(normalize_text(bullet.group(1)))
            set_paragraph_font(paragraph, size=10.5, color=TEXT)
            set_paragraph_spacing(paragraph, after=1.5, line=1.12)
            i += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)", stripped)
        if numbered:
            flush_paragraph(doc, paragraph_lines)
            paragraph = doc.add_paragraph(style="List Number")
            paragraph.add_run(normalize_text(numbered.group(1)))
            set_paragraph_font(paragraph, size=10.5, color=TEXT)
            set_paragraph_spacing(paragraph, after=1.5, line=1.12)
            i += 1
            continue

        paragraph_lines.append(stripped)
        i += 1

    flush_paragraph(doc, paragraph_lines)


def main() -> None:
    if not INPUT_MD.exists():
        raise FileNotFoundError(f"Missing input markdown: {INPUT_MD}")

    output_docx = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else OUTPUT_DOCX
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    markdown = INPUT_MD.read_text(encoding="utf-8")

    doc = Document()
    set_doc_defaults(doc)
    add_footer(doc.sections[0])
    add_cover(doc)
    add_markdown_body(doc, markdown)

    doc.core_properties.title = "CoachLink 商业计划书 v2.1"
    doc.core_properties.subject = "线上自由教练的商业化操作系统"
    doc.core_properties.author = "CoachLink"
    doc.save(output_docx)

    loaded = Document(str(output_docx))
    paragraph_count = len(loaded.paragraphs)
    table_count = len(loaded.tables)
    print(f"Generated {output_docx}")
    print(f"Size: {output_docx.stat().st_size} bytes")
    print(f"Paragraphs: {paragraph_count}")
    print(f"Tables: {table_count}")


if __name__ == "__main__":
    main()
